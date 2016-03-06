import os
import openpyxl


def write(workbook_name, row, col, value):
    if not (str(workbook_name).endswith('.xlsx') or str(workbook_name).endswith('.xls')):
        workbook_name = str(workbook_name) + '.xlsx'
    if os.path.exists(workbook_name):
        workbook = openpyxl.load_workbook(workbook_name)
    else:
        workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.cell(row=row, column=col, value=value)
    workbook.save(filename=workbook_name)
