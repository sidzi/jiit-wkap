import os
import openpyxl


def read(workbook_name):
    if not (str(workbook_name).endswith('.xlsx') or str(workbook_name).endswith('.xls')):
        workbook_name = str(workbook_name) + '.xlsx'
    if os.path.exists(workbook_name):
        workbook = openpyxl.load_workbook(workbook_name)
    else:
        print("Excel File Not Found ! ")
        raise Exception
    worksheet = workbook.active
    readData = [[0 for x in range(worksheet.max_column + 1)] for x in range(worksheet.max_row + 1)]
    j = 1
    for row in worksheet.rows:
        i = 1
        for cell in row:
            if cell.value is None:
                readData[j][i] = ""
            else:
                readData[j][i] = float(cell.value)
            i += 1
        j += 1
    return readData


if __name__ == "__main__":
    read('../testData.xlsx')
